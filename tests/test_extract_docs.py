import docx
import pymupdf

from conftest import local_item


def _pdf(path, text):
    d = pymupdf.open()
    page = d.new_page()
    if text:
        page.insert_text((72, 72), text, fontsize=11)
    d.save(path)
    d.close()


async def test_pdf_text(fake_pipeline, tmp_path):
    p = tmp_path / "memo.pdf"
    _pdf(p, "Vendor memo\n@Jordan Blake reviewed the terms.\nWe decided to accept net-30 terms.")
    r = await fake_pipeline.run(local_item(p))
    body = (fake_pipeline.vault.root / r.note.note_path).read_text(encoding="utf-8")
    assert "pages: 1" in body and "ocr:" not in body
    assert "--- page 1 ---" in body and "net-30" in body
    assert (fake_pipeline.vault.root / "people" / "Jordan Blake.md").exists()
    assert list((fake_pipeline.vault.root / "attachments").glob("*-memo.pdf"))


async def test_pdf_scanned_falls_back_to_llm(fake_pipeline, tmp_path):
    p = tmp_path / "scan.pdf"
    _pdf(p, "")
    r = await fake_pipeline.run(local_item(p))
    body = (fake_pipeline.vault.root / r.note.note_path).read_text(encoding="utf-8")
    assert "ocr: claude" in body and "fake transcription" in body


async def test_docx(fake_pipeline, tmp_path):
    p = tmp_path / "plan.docx"
    d = docx.Document()
    d.core_properties.title = "Hiring plan"
    d.add_paragraph("@Marko Horvat owns the pipeline.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Role", "Count"
    d.save(p)
    r = await fake_pipeline.run(local_item(p))
    body = (fake_pipeline.vault.root / r.note.note_path).read_text(encoding="utf-8")
    assert "title: Hiring plan" in body and "Role\tCount" in body
    assert (fake_pipeline.vault.root / "people" / "Marko Horvat.md").exists()


async def test_image_uses_structurer(fake_pipeline, tmp_path):
    p = tmp_path / "shot.png"
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 8, 8), False)
    pix.save(p)
    r = await fake_pipeline.run(local_item(p))
    body = (fake_pipeline.vault.root / r.note.note_path).read_text(encoding="utf-8")
    assert "fake image description of shot.png" in body
