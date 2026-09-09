"""Document extractors: PDF (PyMuPDF, Claude fallback for scans), DOCX, images (Claude vision)."""

from __future__ import annotations

import asyncio
import logging
from pathlib import Path

from .extract import ExtractorRegistry, UnsupportedTypeError
from .models import Extracted, IngestItem

log = logging.getLogger(__name__)

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
SCANNED_THRESHOLD = 200  # chars; below this a PDF is treated as image-only
PDF_LLM_MAX_MB = 32


class PdfExtractor:
    def __init__(self, structurer=None):
        self.structurer = structurer

    async def extract(self, item: IngestItem, progress) -> Extracted:
        p = Path(item.local_path)  # type: ignore[arg-type]
        pages, text = await asyncio.to_thread(self._read, p)
        meta = {"format": "pdf"}
        if len(text.strip()) < SCANNED_THRESHOLD and self.structurer is not None:
            if p.stat().st_size > PDF_LLM_MAX_MB * 1024 * 1024:
                raise UnsupportedTypeError("scanned PDF too large for vision fallback")
            await progress.update("reading scanned pdf", p.name)
            text = await self.structurer.read_pdf(p)
            meta["ocr"] = "claude"
        return Extracted(item=item, kind="document", text=text, page_count=pages, meta=meta)

    @staticmethod
    def _read(p: Path) -> tuple[int, str]:
        import pymupdf

        with pymupdf.open(p) as doc:
            parts = []
            for i, page in enumerate(doc, 1):
                t = page.get_text("text").strip()
                if t:
                    parts.append(f"--- page {i} ---\n{t}")
            return doc.page_count, "\n\n".join(parts)


class DocxExtractor:
    async def extract(self, item: IngestItem, progress) -> Extracted:
        p = Path(item.local_path)  # type: ignore[arg-type]
        text = await asyncio.to_thread(self._read, p)
        return Extracted(item=item, kind="document", text=text, meta={"format": "docx"})

    @staticmethod
    def _read(p: Path) -> str:
        import docx

        d = docx.Document(str(p))
        parts = [para.text for para in d.paragraphs if para.text.strip()]
        for table in d.tables:
            for row in table.rows:
                parts.append("\t".join(c.text.strip() for c in row.cells))
        core = d.core_properties
        head = []
        if core.title:
            head.append(f"# {core.title}")
        if core.author:
            head.append(f"Author: {core.author}")
        return "\n".join(head + [""] + parts).strip()


class ImageExtractor:
    def __init__(self, structurer):
        self.structurer = structurer

    async def extract(self, item: IngestItem, progress) -> Extracted:
        p = Path(item.local_path)  # type: ignore[arg-type]
        await progress.update("describing image", p.name)
        text = await self.structurer.describe_image(p)
        return Extracted(
            item=item, kind="document", text=text, meta={"format": p.suffix.lstrip(".")}
        )


def register_docs(reg: ExtractorRegistry, *, structurer=None) -> None:
    reg.register({".pdf"}, PdfExtractor(structurer))
    reg.register({".docx"}, DocxExtractor())
    if structurer is not None:
        reg.register(IMAGE_EXT, ImageExtractor(structurer))
